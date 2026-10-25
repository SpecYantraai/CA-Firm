"""Built-in Office editing adapted from the supplied Assessify editor workflow.

Only changed OOXML parts are written back, preserving unrelated package content.
"""
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
import math
import posixpath
import re
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W, "s": S}
MAX_CELLS = 100_000
MAX_NODES = 20_000


class OfficeError(ValueError):
    pass


def xml(data):
    return ET.fromstring(data, ET.XMLParser(resolve_entities=False, no_network=True))


def package(data):
    try:
        archive = ZipFile(BytesIO(data))
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > 150_000_000:
            archive.close()
            raise OfficeError("This Office file is too large for the built-in editor.")
        if len({entry.filename for entry in entries}) != len(entries):
            archive.close()
            raise OfficeError("This Office file contains duplicate package entries.")
        if any(entry.filename.startswith("_xmlsignatures/") or entry.filename.lower().endswith("vbaproject.bin") for entry in entries):
            archive.close()
            raise OfficeError("Digitally signed or macro-enabled Office files cannot be edited with the built-in editor.")
        return archive
    except BadZipFile as exc:
        raise OfficeError("This file is not a valid DOCX or XLSX document.") from exc


def rewrite(archive, updates, removed=()):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in archive.infolist():
            if entry.filename not in removed:
                target.writestr(entry, updates.get(entry.filename, archive.read(entry.filename)))
        existing = set(archive.namelist())
        for name, content in updates.items():
            if name not in existing and name not in removed:
                target.writestr(name, content)
    return output.getvalue()


def encoded(root):
    return ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def word_nodes(root):
    nodes = list(root.find(f"{{{W}}}body").iter())
    return {f"n{index}": node for index, node in enumerate(nodes)}


def word_view(data):
    with package(data) as archive:
        root = xml(archive.read("word/document.xml"))
        if "word/settings.xml" in archive.namelist():
            protection = xml(archive.read("word/settings.xml")).find("w:documentProtection", NS)
            if protection is not None and protection.get(f"{{{W}}}enforcement") in {"1", "true", "on"}:
                raise OfficeError("This Word document is protected. Use an authorized Office editor to unlock it.")
        nodes = word_nodes(root)
        ids = {node: key for key, node in nodes.items()}
        list_formats = {}
        if "word/numbering.xml" in archive.namelist():
            numbering = xml(archive.read("word/numbering.xml"))
            abstracts = {node.get(f"{{{W}}}abstractNumId"): node for node in numbering.findall("w:abstractNum", NS)}
            for num in numbering.findall("w:num", NS):
                abstract_id = num.find("w:abstractNumId", NS)
                abstract = abstracts.get(abstract_id.get(f"{{{W}}}val")) if abstract_id is not None else None
                fmt = abstract.find("w:lvl/w:numFmt", NS) if abstract is not None else None
                if fmt is not None:
                    list_formats[num.get(f"{{{W}}}numId")] = fmt.get(f"{{{W}}}val")
        if len(nodes) > MAX_NODES:
            raise OfficeError("This document is too large for the built-in editor.")

        def read_blocks(elements):
            blocks = []
            for element in elements:
                block = read(element)
                if blocks and block["type"] in {"bulletList", "orderedList"} and blocks[-1]["type"] == block["type"]:
                    blocks[-1]["content"].extend(block["content"])
                else:
                    blocks.append(block)
            return blocks

        def read(node):
            kind = ET.QName(node).localname
            source = {"sourceId": ids[node]}
            if kind == "p":
                # Fields, drawings, hyperlinks and tracked changes must not be flattened.
                safe = node.find("w:pPr/w:sectPr", NS) is None and all(child.tag in {f"{{{W}}}pPr", f"{{{W}}}r"} for child in node)
                safe = safe and all(child.tag in {f"{{{W}}}rPr", f"{{{W}}}t", f"{{{W}}}br", f"{{{W}}}tab"} for run in node.findall("w:r", NS) for child in run)
                if not safe:
                    return protected(node, source, "Embedded content")
                style = node.find("w:pPr/w:pStyle", NS)
                style_id = style.get(f"{{{W}}}val", "") if style is not None else ""
                heading = re.fullmatch(r"Heading([1-6])", style_id, re.I)
                content = []
                for run in node.findall("w:r", NS):
                    marks = []
                    for mark, tag in (("bold", "b"), ("italic", "i"), ("underline", "u"), ("strike", "strike")):
                        prop = run.find(f"w:rPr/w:{tag}", NS)
                        if prop is not None and prop.get(f"{{{W}}}val", "1") not in {"0", "false", "none"}:
                            marks.append({"type": mark})
                    for child in run:
                        if child.tag == f"{{{W}}}br":
                            content.append({"type": "hardBreak"})
                        elif child.tag in {f"{{{W}}}t", f"{{{W}}}tab"}:
                            text = "\t" if child.tag == f"{{{W}}}tab" else child.text or ""
                            if text:
                                content.append({"type": "text", "text": text, "marks": marks})
                paragraph = {"type": "heading" if heading else "paragraph", "attrs": source | ({"level": int(heading[1])} if heading else {}), "content": content}
                number = node.find("w:pPr/w:numPr/w:numId", NS)
                if not heading and (number is not None or style_id in {"ListBullet", "ListNumber"}):
                    fmt = list_formats.get(number.get(f"{{{W}}}val")) if number is not None else "bullet" if style_id == "ListBullet" else "decimal"
                    return {"type": "bulletList" if fmt == "bullet" else "orderedList", "content": [{"type": "listItem", "content": [paragraph]}]}
                return paragraph
            if kind == "tbl":
                # Complex/merged tables remain untouched; simple tables support text editing.
                if node.findall(".//w:vMerge", NS) or node.findall(".//w:gridSpan", NS) or len(node.findall(".//w:tbl", NS)):
                    return protected(node, source, "Complex table")
                rows = []
                for row in node.findall("w:tr", NS):
                    cells = []
                    for cell in row.findall("w:tc", NS):
                        children = read_blocks(child for child in cell if child.tag != f"{{{W}}}tcPr")
                        cells.append({"type": "tableCell", "attrs": {"sourceId": ids[cell]}, "content": children or [{"type": "paragraph"}]})
                    rows.append({"type": "tableRow", "attrs": {"sourceId": ids[row]}, "content": cells})
                return {"type": "table", "attrs": source, "content": rows}
            return protected(node, source, "Document content")

        def protected(node, attrs, label):
            text = " ".join(node.itertext()).strip()[:160]
            return {"type": "preservedBlock", "attrs": attrs | {"label": label, "text": text}}

        content = read_blocks(node for node in root.find("w:body", NS) if node.tag != f"{{{W}}}sectPr")
        return {"type": "doc", "content": content or [{"type": "paragraph"}]}


def write_word(data, document):
    original = word_view(data)
    if document == original:
        return data
    with package(data) as archive:
        root = xml(archive.read("word/document.xml"))
        originals = word_nodes(root)
        original_json = {}
        original_lists = {}
        protected_ids = set()

        def index(node, list_style=None):
            if node["type"] in {"bulletList", "orderedList"}:
                list_style = "ListBullet" if node["type"] == "bulletList" else "ListNumber"
            source = node.get("attrs", {}).get("sourceId")
            if source:
                original_json[source] = node
                original_lists[source] = list_style
                if node["type"] == "preservedBlock":
                    protected_ids.add(source)
            for child in node.get("content", []):
                index(child, list_style)

        index(original)
        seen = set()
        preserved = set()
        budget = [MAX_NODES]
        numbering_updates = {}
        list_ids = {}

        def canonical(value):
            if isinstance(value, list):
                return [canonical(item) for item in value]
            if not isinstance(value, dict):
                return value
            result = {}
            for key, item in value.items():
                if item is None or item == [] or item == {}:
                    continue
                if key == "attrs":
                    item = {name: val for name, val in item.items() if name in {"sourceId", "level", "label", "text"} and val is not None}
                    if not item:
                        continue
                result[key] = canonical(item)
            if "marks" in result:
                result["marks"] = sorted(result["marks"], key=lambda mark: mark["type"])
            return result

        def list_numbering(style):
            if style in list_ids:
                return list_ids[style]
            path = "word/numbering.xml"
            number_root = xml(numbering_updates[path]) if path in numbering_updates else xml(archive.read(path)) if path in archive.namelist() else ET.Element(f"{{{W}}}numbering", nsmap={"w": W})
            abstract_id = max([int(item.get(f"{{{W}}}abstractNumId")) for item in number_root.findall("w:abstractNum", NS)] + [-1]) + 1
            number_id = max([int(item.get(f"{{{W}}}numId")) for item in number_root.findall("w:num", NS)] + [0]) + 1
            abstract = ET.SubElement(number_root, f"{{{W}}}abstractNum", {f"{{{W}}}abstractNumId": str(abstract_id)})
            level = ET.SubElement(abstract, f"{{{W}}}lvl", {f"{{{W}}}ilvl": "0"})
            for tag, val in (("start", "1"), ("numFmt", "bullet" if style == "ListBullet" else "decimal"), ("lvlText", "\u2022" if style == "ListBullet" else "%1.")):
                ET.SubElement(level, f"{{{W}}}{tag}", {f"{{{W}}}val": val})
            num = ET.SubElement(number_root, f"{{{W}}}num", {f"{{{W}}}numId": str(number_id)})
            ET.SubElement(num, f"{{{W}}}abstractNumId", {f"{{{W}}}val": str(abstract_id)})
            numbering_updates[path] = encoded(number_root)
            if path not in archive.namelist():
                relpath = "word/_rels/document.xml.rels"
                rels = xml(archive.read(relpath))
                rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
                rel_id = "rIdNativeNumbering"
                if not any(item.get("Type", "").endswith("/numbering") for item in rels):
                    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", Id=rel_id, Type=R + "/numbering", Target="numbering.xml")
                numbering_updates[relpath] = encoded(rels)
                types = xml(archive.read("[Content_Types].xml"))
                ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/word/numbering.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml")
                numbering_updates["[Content_Types].xml"] = encoded(types)
            list_ids[style] = number_id
            return number_id

        def build(node, depth=0, list_style=None):
            budget[0] -= 1
            if budget[0] < 0 or depth > 30 or not isinstance(node, dict):
                raise OfficeError("The edited document is too complex.")
            kind = node.get("type")
            attrs = node.get("attrs") or {}
            source = attrs.get("sourceId")
            old = originals.get(source)
            if source in seen:
                old = None  # Splitting a paragraph can copy its source attribute.
            if source:
                seen.add(source)
            if kind == "preservedBlock":
                if source not in protected_ids or source in preserved:
                    raise OfficeError("An embedded document block was changed. Reload the document.")
                preserved.add(source)
                return [deepcopy(originals[source])]
            if old is not None and canonical(original_json.get(source)) == canonical(node) and original_lists.get(source) == list_style:
                # Includes any protected descendants in an unchanged table.
                def mark(child):
                    child_source = (child.get("attrs") or {}).get("sourceId")
                    if child_source in protected_ids:
                        preserved.add(child_source)
                    for item in child.get("content", []):
                        mark(item)
                mark(node)
                return [deepcopy(old)]
            if kind in {"bulletList", "orderedList", "listItem", "blockquote"}:
                style = "ListBullet" if kind == "bulletList" else "ListNumber" if kind == "orderedList" else list_style
                return [element for child in node.get("content", []) for element in build(child, depth + 1, style)]
            tags = {"paragraph": "p", "heading": "p", "table": "tbl", "tableRow": "tr", "tableCell": "tc", "tableHeader": "tc"}
            if kind not in tags:
                raise OfficeError(f"Unsupported document block: {kind}")
            tag = tags[kind]
            result = ET.Element(f"{{{W}}}{tag}")
            prop_tag = {"p": "pPr", "tbl": "tblPr", "tr": "trPr", "tc": "tcPr"}[tag]
            if old is not None and old.tag == result.tag:
                prop = old.find(f"{{{W}}}{prop_tag}")
                if prop is not None:
                    result.append(deepcopy(prop))
                if tag == "tbl":
                    grid = old.find("w:tblGrid", NS)
                    if grid is not None:
                        result.append(deepcopy(grid))
            if tag == "p":
                ppr = result.find("w:pPr", NS)
                if ppr is None:
                    ppr = ET.SubElement(result, f"{{{W}}}pPr")
                style = f"Heading{max(1, min(6, int(attrs.get('level', 1))))}" if kind == "heading" else list_style
                if style:
                    element = ppr.find("w:pStyle", NS)
                    if element is None:
                        element = ET.SubElement(ppr, f"{{{W}}}pStyle")
                    element.set(f"{{{W}}}val", style)
                if list_style and not (old is not None and original_lists.get(source) == list_style and old.find("w:pPr/w:numPr", NS) is not None):
                    for item in list(ppr.findall("w:numPr", NS)):
                        ppr.remove(item)
                    numpr = ET.SubElement(ppr, f"{{{W}}}numPr")
                    ET.SubElement(numpr, f"{{{W}}}ilvl", {f"{{{W}}}val": "0"})
                    ET.SubElement(numpr, f"{{{W}}}numId", {f"{{{W}}}val": str(list_numbering(list_style))})
                elif kind == "paragraph" and old is not None and original_json.get(source, {}).get("type") == "heading":
                    for element in list(ppr.findall("w:pStyle", NS)):
                        ppr.remove(element)
                if not list_style and original_lists.get(source):
                    for element in list(ppr.findall("w:numPr", NS)) + list(ppr.findall("w:pStyle", NS)):
                        ppr.remove(element)
                base = old.find("w:r/w:rPr", NS) if old is not None else None
                for child in node.get("content", []):
                    run = ET.SubElement(result, f"{{{W}}}r")
                    if child.get("type") == "hardBreak":
                        ET.SubElement(run, f"{{{W}}}br")
                        continue
                    if child.get("type") != "text" or not isinstance(child.get("text"), str):
                        raise OfficeError("Unsupported inline document content.")
                    rpr = deepcopy(base) if base is not None else ET.Element(f"{{{W}}}rPr")
                    marks = {mark["type"] for mark in child.get("marks", [])}
                    for mark, flag in (("bold", "b"), ("italic", "i"), ("underline", "u"), ("strike", "strike")):
                        existing_flags = list(rpr.findall(f"w:{flag}", NS))
                        for element in existing_flags:
                            rpr.remove(element)
                        if mark in marks or existing_flags:
                            flag_node = ET.SubElement(rpr, f"{{{W}}}{flag}")
                            flag_node.set(f"{{{W}}}val", ("single" if flag == "u" else "1") if mark in marks else ("none" if flag == "u" else "0"))
                    run.append(rpr)
                    text = ET.SubElement(run, f"{{{W}}}t")
                    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    text.text = child["text"]
            else:
                for child in node.get("content", []):
                    result.extend(build(child, depth + 1))
                if tag == "tbl":
                    rows = result.findall("w:tr", NS)
                    columns = max([len(row.findall("w:tc", NS)) for row in rows] + [1])
                    grid = result.find("w:tblGrid", NS)
                    if grid is None or len(grid) != columns:
                        if grid is not None:
                            result.remove(grid)
                        grid = ET.Element(f"{{{W}}}tblGrid")
                        for _ in range(columns):
                            ET.SubElement(grid, f"{{{W}}}gridCol", {f"{{{W}}}w": str(9000 // columns)})
                        result.insert(1 if result.find("w:tblPr", NS) is not None else 0, grid)
            return [result]

        if document.get("type") != "doc":
            raise OfficeError("A document body is required.")
        replacement = [element for child in document.get("content", []) for element in build(child)]
        if preserved != protected_ids:
            raise OfficeError("Embedded content cannot be removed in the built-in editor. Undo the deletion or use an external Office editor.")
        body = root.find("w:body", NS)
        sections = [deepcopy(child) for child in body if child.tag == f"{{{W}}}sectPr"]
        body[:] = replacement + sections
        return rewrite(archive, {"word/document.xml": encoded(root), **numbering_updates})


def safe_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def check_workbook(archive):
    count = 0
    for name in sheet_paths(archive).values():
        root = xml(archive.read(name))
        count += len(root.findall(".//s:sheetData/s:row/s:c", NS))
        if count > MAX_CELLS:
            raise OfficeError("This workbook exceeds the built-in editor's 100,000-cell limit. Use the external editor.")
        for merged in root.findall("s:mergeCells/s:mergeCell", NS):
            left, top, right, bottom = range_boundaries(merged.get("ref"))
            if (right - left + 1) * (bottom - top + 1) > MAX_CELLS:
                raise OfficeError("This workbook has a merged range too large for the built-in editor.")


def array_member(ranges, address):
    row, col = coordinate_to_tuple(address)
    for bounds in ranges:
        left, top, right, bottom = range_boundaries(bounds)
        if left <= col <= right and top <= row <= bottom:
            return True
    return False


def workbook_view(data):
    with package(data) as archive:
        check_workbook(archive)
    workbook = load_workbook(BytesIO(data), data_only=False)
    cached = load_workbook(BytesIO(data), data_only=True)
    try:
        sheets = []
        count = 0
        for sheet in workbook.worksheets:
            cells = {}
            arrays = list(sheet.array_formulae.values())
            # Sparse iteration avoids materializing a million-row formatted sheet.
            for cell in list(sheet._cells.values()):
                if cell.value is None and not cell.has_style:
                    continue
                count += 1
                if count > MAX_CELLS:
                    raise OfficeError("This workbook exceeds the built-in editor's 100,000-cell limit. Use the external editor.")
                formula = cell.data_type == "f"
                complex_formula = formula and not isinstance(cell.value, str)
                cells[cell.coordinate] = {"raw": "Array formula (preserved)" if complex_formula else safe_value(cell.value), "formula": formula,
                    "cached": safe_value(cached[sheet.title][cell.coordinate].value) if formula else None,
                    "editable": not complex_formula and not array_member(arrays, cell.coordinate) and not (sheet.protection.sheet and cell.protection.locked),
                    "bold": bool(cell.font.bold), "italic": bool(cell.font.italic),
                    "number_format": cell.number_format,
                    "background": cell.fill.fgColor.rgb if cell.fill.patternType == "solid" and cell.fill.fgColor.type == "rgb" else None}
            sheets.append({"name": sheet.title, "cells": cells, "n_rows": min(1048576, max(50, sheet.max_row)),
                "n_cols": min(16384, max(12, sheet.max_column)), "protected": bool(sheet.protection.sheet),
                "merged": [str(item) for item in sheet.merged_cells.ranges], "protected_ranges": arrays, "hidden": sheet.sheet_state != "visible"})
        return {"sheets": sheets}
    finally:
        workbook.close()
        cached.close()


def sheet_paths(archive):
    workbook = xml(archive.read("xl/workbook.xml"))
    rels = xml(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.get("Id"): item.get("Target") for item in rels}
    return {item.get("name"): posixpath.normpath("xl/" + targets[item.get(f"{{{R}}}id")]) if not targets[item.get(f"{{{R}}}id")].startswith("/") else targets[item.get(f"{{{R}}}id")].lstrip("/") for item in workbook.find("s:sheets", NS)}


def coerce(value):
    if not isinstance(value, str) or len(value) > 32767:
        raise OfficeError("Each cell must contain at most 32,767 characters.")
    if value == "":
        return None, None
    if value.startswith("'"):
        return "inlineStr", value[1:]
    if value.startswith("="):
        return "f", value[1:]
    if value.lower() in {"true", "false"}:
        return "b", "1" if value.lower() == "true" else "0"
    if re.fullmatch(r"[+-]?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?", value) and len(re.sub(r"[^0-9]", "", value)) <= 15:
        if math.isfinite(float(value)):
            return "n", value
    return "inlineStr", value


def write_workbook(data, edits):
    with package(data) as archive:
        check_workbook(archive)
    workbook = load_workbook(BytesIO(data), data_only=False)
    try:
        with package(data) as archive:
            paths = sheet_paths(archive)
            updates = {}
            count = 0
            for name, changes in edits.items():
                if name not in paths or name not in workbook.sheetnames:
                    raise OfficeError("The worksheet does not exist.")
                sheet = workbook[name]
                arrays = list(sheet.array_formulae.values())
                root = xml(archive.read(paths[name]))
                sheet_data = root.find("s:sheetData", NS)
                rows = {int(row.get("r")): row for row in sheet_data}
                for address, value in changes.items():
                    count += 1
                    if count > 10_000 or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", address):
                        raise OfficeError("Invalid cell address or too many edited cells.")
                    row_no, col_no = coordinate_to_tuple(address)
                    if row_no > 1048576 or col_no > 16384:
                        raise OfficeError("Cell address is outside Excel's worksheet limits.")
                    cell = sheet[address]
                    if cell.__class__.__name__ == "MergedCell" or (sheet.protection.sheet and cell.protection.locked):
                        raise OfficeError(f"Cell {name}!{address} is protected or part of a merged cell.")
                    if array_member(arrays, address):
                        raise OfficeError("Array formula ranges are protected in the built-in editor.")
                    if cell.data_type == "f" and not isinstance(cell.value, str):
                        raise OfficeError("Array and data-table formulas are protected in the built-in editor.")
                    row = rows.get(row_no)
                    if row is None:
                        row = ET.Element(f"{{{S}}}row", r=str(row_no))
                        following = next((item for item in sheet_data if int(item.get("r")) > row_no), None)
                        if following is not None:
                            following.addprevious(row)
                        else:
                            sheet_data.append(row)
                        rows[row_no] = row
                    element = next((item for item in row if item.get("r") == address), None)
                    if element is None:
                        element = ET.Element(f"{{{S}}}c", r=address)
                        following = next((item for item in row if coordinate_to_tuple(item.get("r"))[1] > col_no), None)
                        if following is not None:
                            following.addprevious(element)
                        else:
                            row.append(element)
                    for child in list(element):
                        if child.tag in {f"{{{S}}}f", f"{{{S}}}v", f"{{{S}}}is"}:
                            element.remove(child)
                    element.attrib.pop("t", None)
                    kind, text = coerce(value)
                    if kind == "inlineStr":
                        element.set("t", kind)
                        inline = ET.SubElement(element, f"{{{S}}}is")
                        node = ET.SubElement(inline, f"{{{S}}}t")
                        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        node.text = text
                    elif kind:
                        if kind != "f":
                            element.set("t", kind)
                        ET.SubElement(element, f"{{{S}}}{'f' if kind == 'f' else 'v'}").text = text
                # Existing dimensions may no longer include newly edited cells.
                dimension = root.find("s:dimension", NS)
                if dimension is not None:
                    root.remove(dimension)
                updates[paths[name]] = encoded(root)
            for path in paths.values():
                root = xml(updates.get(path, archive.read(path)))
                for cell in root.findall(".//s:c", NS):
                    formula = cell.find("s:f", NS)
                    if formula is not None:
                        if formula.get("t") == "shared":
                            sheet_name = next(name for name, candidate in paths.items() if candidate == path)
                            expanded = workbook[sheet_name][cell.get("r")].value
                            if isinstance(expanded, str) and expanded.startswith("="):
                                formula.attrib.clear()
                                formula.text = expanded[1:]
                        for cache in list(cell.findall("s:v", NS)):
                            cell.remove(cache)
                updates[path] = encoded(root)
            root = xml(archive.read("xl/workbook.xml"))
            calc = root.find("s:calcPr", NS)
            if calc is None:
                calc = ET.SubElement(root, f"{{{S}}}calcPr")
            calc.attrib.update({"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})
            updates["xl/workbook.xml"] = encoded(root)
            for path in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                root = xml(archive.read(path))
                for child in list(root):
                    if "calcChain" in (child.get("Type", "") + child.get("PartName", "")):
                        root.remove(child)
                updates[path] = encoded(root)
            return rewrite(archive, updates, removed={"xl/calcChain.xml"})
    finally:
        workbook.close()
