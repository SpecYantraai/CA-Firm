"""Office round trips and authenticated/versioned working-paper saves."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.database import Base, get_db
from app.core.deps import get_current_user
from app.models.models import User, Engagement, EngagementUser, Section, Folder, WorkingPaper, FileVersion
from app.services.native_office import OfficeError, word_view, workbook_view, write_word, write_workbook, xml, NS, W
from app.api.v1.endpoints.native_editor import router
from app.api.v1.endpoints.wps import router as wp_router


def word_bytes(embedded=False):
    document = Document()
    document.add_heading("Audit conclusion", 1)
    paragraph = document.add_paragraph("Revenue is ")
    paragraph.add_run("complete").bold = True
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Evidence"
    document.sections[0].header.paragraphs[0].text = "Confidential audit working paper"
    if embedded:
        field = OxmlElement("w:fldSimple")
        paragraph._p.append(field)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def excel_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trial balance"
    sheet["A1"] = "Account"
    sheet["A2"] = "00123"
    sheet["B2"] = 10
    sheet["B3"] = 20
    sheet["B4"] = "=SUM(B2:B3)"
    sheet["B2"].number_format = '#,##0.00'
    sheet["D1"] = "Merged"
    sheet.merge_cells("D1:E1")
    workbook.create_sheet("Summary")["A1"] = "='Trial balance'!B4"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class OfficeRoundTripTests(unittest.TestCase):
    def test_word_edit_preserves_tables_headers_and_other_parts(self):
        before = word_bytes()
        view = word_view(before)
        view["content"][1]["content"][1]["text"] = "accurate"
        view["content"][1]["content"][1]["marks"] = [{"type": "underline"}]
        table_p = view["content"][2]["content"][0]["content"][0]["content"][0]
        table_p["content"][0]["text"] = "Linked evidence"
        after = write_word(before, view)
        document = Document(BytesIO(after))
        self.assertEqual(document.paragraphs[1].text, "Revenue is accurate")
        self.assertTrue(document.paragraphs[1].runs[-1].underline)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Linked evidence")
        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
        with ZipFile(BytesIO(before)) as original, ZipFile(BytesIO(after)) as saved:
            for name in original.namelist():
                if name != "word/document.xml":
                    self.assertEqual(original.read(name), saved.read(name), name)

    def test_word_protected_blocks_and_new_lists(self):
        before = word_bytes(embedded=True)
        view = word_view(before)
        self.assertEqual(view["content"][1]["type"], "preservedBlock")
        changed = deepcopy(view)
        changed["content"].pop(1)
        with self.assertRaises(OfficeError):
            write_word(before, changed)
        view["content"].append({"type": "bulletList", "content": [{"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "New procedure"}]}]}]})
        after = write_word(before, view)
        document = Document(BytesIO(after))
        self.assertEqual(document.paragraphs[-1].text, "New procedure")
        self.assertIsNotNone(document.paragraphs[-1]._p.find("w:pPr/w:numPr", NS))
        self.assertEqual(word_view(after)["content"][-1]["type"], "bulletList")
        with ZipFile(BytesIO(after)) as archive:
            self.assertIsNotNone(xml(archive.read("word/document.xml")).find(".//w:fldSimple", NS))

    def test_word_no_changes_and_tiptap_default_attributes(self):
        before = word_bytes()
        view = word_view(before)
        self.assertEqual(write_word(before, view), before)
        # ProseMirror normalizes empty attributes/marks. Untouched paragraphs must retain their runs.
        view["content"][0]["content"][0].pop("marks", None)
        view["content"][2]["content"][0]["content"][0]["attrs"].update(colspan=1, rowspan=1, colwidth=None)
        after = write_word(before, view)
        with ZipFile(BytesIO(before)) as original, ZipFile(BytesIO(after)) as saved:
            self.assertEqual(xml(original.read("word/document.xml")).find("w:body/w:p", NS)[1].tag, f"{{{W}}}r")
            self.assertEqual(original.read("word/styles.xml"), saved.read("word/styles.xml"))

    def test_excel_edits_preserve_formula_format_merges_and_extra_parts(self):
        before = excel_bytes()
        container = BytesIO(before)
        with ZipFile(container, "a") as archive:
            archive.writestr("customXml/audit.xml", "<audit>preserve me</audit>")
        before = container.getvalue()
        view = workbook_view(before)
        self.assertEqual(view["sheets"][0]["cells"]["B4"]["raw"], "=SUM(B2:B3)")
        after = write_workbook(before, {"Trial balance": {"B2": "35", "A2": "00123", "C2": "=B2*2", "Z250": "Added", "A3": "'=literal"}})
        workbook = load_workbook(BytesIO(after), data_only=False)
        self.assertEqual(workbook["Trial balance"]["B2"].value, 35)
        self.assertEqual(workbook["Trial balance"]["B2"].number_format, '#,##0.00')
        self.assertEqual(workbook["Trial balance"]["B4"].value, "=SUM(B2:B3)")
        self.assertEqual(workbook["Trial balance"]["C2"].value, "=B2*2")
        self.assertEqual(workbook["Trial balance"]["A2"].value, "00123")
        self.assertEqual(workbook["Trial balance"]["A3"].value, "=literal")
        self.assertEqual(workbook["Trial balance"]["A3"].data_type, "s")
        self.assertEqual(workbook["Trial balance"]["Z250"].value, "Added")
        self.assertIn("D1:E1", str(workbook["Trial balance"].merged_cells))
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        workbook.close()
        with ZipFile(BytesIO(after)) as archive:
            self.assertEqual(archive.read("customXml/audit.xml"), b"<audit>preserve me</audit>")

    def test_excel_validation_and_protection(self):
        before = excel_bytes()
        for edits in ({"Missing": {"A1": "1"}}, {"Trial balance": {"XFE1": "1"}}, {"Trial balance": {"A0": "1"}}, {"Trial balance": {"E1": "1"}}):
            with self.subTest(edits=edits), self.assertRaises(OfficeError):
                write_workbook(before, edits)
        workbook = load_workbook(BytesIO(before))
        workbook["Trial balance"].protection.sheet = True
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaises(OfficeError):
            write_workbook(output.getvalue(), {"Trial balance": {"B2": "1"}})

    def test_array_formula_ranges_are_read_only(self):
        from openpyxl.worksheet.formula import ArrayFormula
        workbook = load_workbook(BytesIO(excel_bytes()))
        workbook["Trial balance"]["F1"] = ArrayFormula(ref="F1:F2", text="=B2:B3*2")
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        data = stream.getvalue()
        self.assertFalse(workbook_view(data)["sheets"][0]["cells"]["F1"]["editable"])
        with self.assertRaises(OfficeError):
            write_workbook(data, {"Trial balance": {"F2": "1"}})


class NativeEditorAPITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_storage = settings.FILE_STORAGE_PATH
        settings.FILE_STORAGE_PATH = self.temp.name
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.Session() as db:
            self.admin = User(user_id="admin", full_name="Admin", email="admin@test.invalid", role="Admin", hashed_password="test")
            self.client = User(user_id="client", full_name="Client", email="client@test.invalid", role="Articled Assistant", hashed_password="test")
            self.other = User(user_id="other", full_name="Other", email="other@test.invalid", role="Articled Assistant", hashed_password="test")
            self.eqcr = User(user_id="eqcr", full_name="EQCR", email="eqcr@test.invalid", role="EQCR Reviewer", hashed_password="test")
            db.add_all([self.admin, self.client, self.other, self.eqcr])
            db.add(Engagement(engagement_id="eng", client_name="Client", financial_year="2026-27", created_by="admin", workflow_override=True))
            db.add(EngagementUser(engagement_id="eng", user_id="client", role="Preparer"))
            db.add(Section(section_id="sec", engagement_id="eng", section_code="2000", section_name="Planning"))
            db.add(Folder(folder_id="folder", engagement_id="eng", section_id="sec", folder_name="Working papers"))
            for ext, data in (("docx", word_bytes()), ("xlsx", excel_bytes())):
                path = Path(self.temp.name) / f"test.{ext}"
                path.write_bytes(data)
                db.add(WorkingPaper(wp_id=ext, engagement_id="eng", section_id="sec", folder_id="folder", wp_number=ext, filename=f"test.{ext}", file_format=ext, file_storage_path=str(path), file_size_bytes=len(data), current_version=1))
            db.commit()
        self.actor = self.client
        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        app.include_router(wp_router, prefix="/api/v1")
        def database():
            with self.Session() as db:
                yield db
        app.dependency_overrides[get_db] = database
        app.dependency_overrides[get_current_user] = lambda: self.actor
        self.api = TestClient(app)

    def tearDown(self):
        self.api.close()
        self.engine.dispose()
        settings.FILE_STORAGE_PATH = self.old_storage
        self.temp.cleanup()

    def test_save_versions_conflicts_and_noop(self):
        path = "/api/v1/wps/docx/native-editor"
        view = self.api.get(path).json()
        self.assertTrue(view["editable"])
        response = self.api.put(path, json={"version": 1, "document": view["document"]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["data"]["current_version"], 1)
        view["document"]["content"][0]["content"][0]["text"] = "Revised audit conclusion"
        response = self.api.put(path, json={"version": 1, "document": view["document"]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["data"]["current_version"], 2)
        with self.Session() as db:
            self.assertEqual(db.query(FileVersion).filter_by(wp_id="docx").count(), 1)
        self.assertEqual(self.api.put(path, json={"version": 1, "document": view["document"]}).status_code, 409)
        response = self.api.put("/api/v1/wps/xlsx/native-editor", json={"version": 1, "edits": {"Trial balance": {"B2": "45"}}})
        self.assertEqual(response.status_code, 200, response.text)

    def test_scope_roles_and_locks(self):
        path = "/api/v1/wps/xlsx/native-editor"
        self.actor = self.other
        self.assertEqual(self.api.get(path).status_code, 403)
        self.assertEqual(self.api.get("/api/v1/wps/xlsx/editor-config").status_code, 403)
        self.actor = self.eqcr
        self.assertFalse(self.api.get(path).json()["editable"])
        self.assertEqual(self.api.put(path, json={"version": 1, "edits": {}}).status_code, 403)
        self.actor = self.client
        for state in ("Archived", "Finalised"):
            with self.Session() as db:
                db.get(Engagement, "eng").status = state
                db.commit()
            self.assertFalse(self.api.get(path).json()["editable"])
            self.assertEqual(self.api.put(path, json={"version": 1, "edits": {}}).status_code, 403)
        with self.Session() as db:
            engagement = db.get(Engagement, "eng")
            engagement.status = "Active"
            engagement.workflow_override = False
            db.commit()
        self.assertFalse(self.api.get(path).json()["editable"])


if __name__ == "__main__":
    unittest.main()
