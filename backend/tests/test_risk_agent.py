"""Run from backend: python -m unittest discover -s tests -p test_risk_agent.py."""
from copy import deepcopy
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.core.deps import get_current_user
from app.core.config import settings
from app.models.models import User, Engagement, EngagementUser, Section, WorkingPaper, RiskAssessmentSession
from app.services import risk_agent as bridge
from app.api.v1.endpoints.risk_agent import router

sys.path.insert(0, str(bridge.VENDOR))
from agents.risk_assessment_agent import RiskAssessmentAgent
from utils import knowledge_base as kb
from utils.materiality import calculate_materiality, calculate_materiality_flexible
from utils.thresholds import calculate_thresholds
from utils.financial_parser import parse_financial_statements, get_mcotabd_candidates


def workbook_bytes(account="Inventories"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "BS"
    sheet.cell(4, 4, "Current year")
    sheet.cell(5, 2, account)
    sheet.cell(5, 3, "6.1")
    sheet.cell(5, 4, 125)
    sheet.cell(5, 5, 999999)  # Prior year must not replace current year.
    sheet.cell(6, 2, "TOTAL ASSETS")
    sheet.cell(6, 4, 500)
    pl = workbook.create_sheet("P&L")
    pl.cell(7, 2, "Revenue from operations")
    pl.cell(7, 3, "6.18")
    pl.cell(7, 4, 1200)
    pl.cell(20, 2, "Profit/(loss) before tax")
    pl.cell(20, 4, 100)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class StubLLM:
    def complete_json(self, **kwargs):
        return []


class AgentIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_storage = settings.FILE_STORAGE_PATH
        settings.FILE_STORAGE_PATH = self.temp.name
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.Session() as db:
            self.client_user = User(user_id="client", full_name="Client", email="client@example.test", hashed_password="test", role="Articled Assistant", initials="CL")
            self.admin = User(user_id="admin", full_name="Firm Admin", email="admin@example.test", hashed_password="test", role="Admin", initials="FA")
            db.add_all([self.client_user, self.admin])
            db.add_all([Engagement(engagement_id=key, client_name=name, financial_year="2025-26", created_by="admin", workflow_override=True) for key, name in (("eng-a", "Client A"), ("eng-b", "Client B"))])
            db.add(EngagementUser(engagement_id="eng-a", user_id="client", role="Preparer"))
            db.add_all([Section(section_id=f"sec-{key}", engagement_id=key, section_code="2000", section_name="Planning") for key in ("eng-a", "eng-b")])
            db.commit()
        self.actor = self.client_user
        app = FastAPI()
        app.include_router(router, prefix="/api/v1")
        def database():
            with self.Session() as db:
                yield db
        app.dependency_overrides[get_db] = database
        app.dependency_overrides[get_current_user] = lambda: self.actor
        self.api = TestClient(app)
        self.url = "/api/v1/engagements/eng-a/risk-agent"
        self.factory_patch = patch("app.core.database.SessionLocal", self.Session)
        self.factory_patch.start()

    def tearDown(self):
        self.factory_patch.stop()
        self.api.close()
        self.engine.dispose()
        settings.FILE_STORAGE_PATH = self.old_storage
        self.temp.cleanup()

    def save(self, operation, data, expected=200):
        revision = self.api.get(self.url).json()["revision"]
        response = self.api.post(f"{self.url}/{operation}", json={"revision": revision, **data})
        self.assertEqual(response.status_code, expected, response.text)
        return response.json()

    def details(self):
        return self.save("details", {"legal_status": "Private Limited Company", "client_type": "New client", "industry": "Manufacturing", "products": "Components", "auditor_risk_notes": "Prior year inventory risk"})

    def test_original_files_are_unchanged(self):
        manifest = json.loads((bridge.BACKEND / "agent_bridge" / "source-sha256.json").read_text(encoding="utf-8-sig"))
        self.assertEqual(len(manifest), 19)
        for name, expected in manifest.items():
            with self.subTest(file=name):
                self.assertEqual(hashlib.sha256((bridge.VENDOR / name).read_bytes()).hexdigest(), expected)
        self.assertFalse((bridge.VENDOR / ".env").exists())

    def test_calculation_and_parser_parity(self):
        figures = {"turnover": 120000000, "total_assets": 50000000, "pbt": 10000000}
        result = bridge.call_agent("materiality", figures)
        original = calculate_materiality(**figures)
        self.assertEqual(result["computed_figure"], original.suggested_figure)
        custom = bridge.call_agent("materiality", {**figures, "benchmark": "Turnover", "percentage": 0.8})
        self.assertEqual(custom["computed_figure"], calculate_materiality_flexible(**figures, ca_benchmark="Turnover", ca_pct=0.8).suggested_figure)
        threshold = bridge.call_agent("thresholds", {"raw_materiality": 1347832, "pm_pct": 80, "ctm_pct": 3})
        expected = calculate_thresholds(1347832, 80, 3)
        self.assertEqual(threshold["rounded_materiality"], expected.rounded_materiality)
        self.assertEqual(threshold["performance_materiality"], expected.performance_materiality)
        self.assertEqual(bridge.call_agent("tvt", {"pm": 1200000, "counts": {"Inventory": 12}}), {"Inventory": 100000})
        path = Path(self.temp.name) / "statements.xlsx"
        path.write_bytes(workbook_bytes())
        parsed = bridge.call_agent("parse", {"document_paths": [str(path)]})
        original_fs = parse_financial_statements(path)
        self.assertEqual(parsed["financials"]["turnover"], original_fs.turnover * 100000)
        self.assertEqual(parsed["parsed_fs"]["bs_items"][0]["value"], 125)
        scope = bridge.call_agent("scope", {**parsed, "thresholds": {"pm": 100}, "materiality": {}})
        self.assertEqual(scope, get_mcotabd_candidates(original_fs, 100))

    def test_client_scope_revision_and_locks(self):
        self.assertEqual(self.api.get("/api/v1/engagements/eng-b/risk-agent").status_code, 403)
        saved = self.details()
        self.assertEqual(saved["state"]["context"]["client_name"], "Client A")
        self.assertEqual(self.api.post(self.url + "/reset", json={"revision": 0}).status_code, 409)
        with self.Session() as db:
            eng = db.get(Engagement, "eng-a")
            eng.status = "Archived"
            db.commit()
        self.assertEqual(self.api.post(self.url + "/reset", json={"revision": saved["revision"]}).status_code, 400)
        self.assertEqual(self.api.get(self.url).status_code, 200)

    def test_invalid_files_and_empty_justification(self):
        self.details()
        revision = self.api.get(self.url).json()["revision"]
        response = self.api.post(self.url + "/financials", data={"revision": revision}, files=[("files", ("bad.xlsx", b"invalid", "application/octet-stream"))])
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.api.get(self.url).json()["revision"], revision)
        response = self.api.post(self.url + "/financials", data={"revision": revision}, files=[("files", ("fs.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))])
        self.assertEqual(response.status_code, 200, response.text)
        self.save("materiality", {"turnover": 120000000, "total_assets": 50000000, "pbt": 10000000, "justification": ""}, expected=400)

    def test_full_workflow_real_agent_with_stub_llm(self):
        raw_rows, *_ = kb._load_raw()
        match = next(row for row in raw_rows if row["industry"] and row["mcotabd"] and row["risk"])
        self.save("details", {"legal_status": "Private Limited Company", "industry": match["industry"]})
        revision = self.api.get(self.url).json()["revision"]
        response = self.api.post(self.url + "/financials", data={"revision": revision}, files=[("files", ("../../statements.xlsx", workbook_bytes(match["mcotabd"]), "application/octet-stream"))])
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()
        self.assertEqual(saved["state"]["documents"][0]["name"], "statements.xlsx")
        self.assertNotIn("path", saved["state"]["documents"][0])
        self.assertNotIn("document_paths", saved["state"]["context"])
        self.save("materiality", {"turnover": 120000000, "total_assets": 50000000, "pbt": 10000000, "justification": "Revenue benchmark selected after review."})
        self.save("thresholds", {"pm_pct": 75, "ctm_pct": 5})
        self.save("scope", {"confirmed_mcotabds": [match["mcotabd"]], "tvt_transaction_counts": {match["mcotabd"]: 12}})
        original_call = bridge.call_agent
        def transport(operation, data=None, **kwargs):
            if operation == "run":
                return RiskAssessmentAgent(StubLLM()).run(deepcopy(data))
            return original_call(operation, data, **kwargs)
        with patch.dict(os.environ, {"NVIDIA_API_KEY": "local-test-no-network"}), patch.object(bridge, "call_agent", side_effect=transport):
            self.save("run", {}, expected=202)
        complete = self.api.get(self.url).json()
        self.assertEqual(complete["state"]["stage"], "complete", complete)
        rows = [{key: row[key] for key in ("key", "assertion", "risk", "controls", "procedures")} for row in complete["state"]["review_rows"]]
        self.assertTrue(rows)
        rows[0]["risk"] = "CA-reviewed risk description"
        reviewed = self.save("review", {"rows": rows})
        self.assertEqual(reviewed["state"]["review_rows"][0]["risk"], rows[0]["risk"])
        self.save("approve", {"rows": rows}, expected=403)
        self.actor = self.admin
        with self.Session() as db:
            engagement = db.get(Engagement, "eng-a")
            engagement.workflow_override = False
            db.commit()
        self.save("approve", {"rows": rows}, expected=400)
        with self.Session() as db:
            db.get(Engagement, "eng-a").workflow_override = True
            db.commit()
        approved = self.save("approve", {"rows": rows})
        self.assertEqual(approved["state"]["stage"], "approved")
        exported = self.api.get(self.url + "/export")
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(exported.json()[0]["risk"], rows[0]["risk"])
        with self.Session() as db:
            wp = db.query(WorkingPaper).one()
            self.assertEqual(wp.section_id, "sec-eng-a")
            self.assertTrue(Path(wp.file_storage_path).is_file())
            self.assertNotIn("document_paths", json.loads(Path(wp.file_storage_path).read_text(encoding="utf-8")))
        self.save("reset", {})
        with self.Session() as db:
            self.assertEqual(db.query(WorkingPaper).count(), 1)
            self.assertEqual(db.get(RiskAssessmentSession, "eng-a").state["stage"], "draft")


if __name__ == "__main__":
    unittest.main()
